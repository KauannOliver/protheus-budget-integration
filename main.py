from __future__ import annotations

import os

import json
import logging
import math
import re
import shutil
import sys
import threading
import unicodedata
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from queue import Empty, Queue
from typing import Any
from zipfile import BadZipFile, ZipFile

import pythoncom
import requests
import win32com.client as win32
from openpyxl import load_workbook
from openpyxl.utils import range_boundaries

from tkinter import (
    Tk,
    BOTH,
    END,
    DISABLED,
    NORMAL,
    EXTENDED,
    StringVar,
    Toplevel,
    Listbox,
    messagebox,
    filedialog,
)
from tkinter import ttk
from tkinter.scrolledtext import ScrolledText


APP_TITLE = "Envio AK2 GN para Protheus"

DEFAULT_WORKBOOK = Path(os.environ.get("PROTHEUS_WORKBOOK", "input/budget.xlsx"))
DEFAULT_EXTRA_WORKBOOK = Path(os.environ.get("PROTHEUS_EXTRA_WORKBOOK", "input/extra-budget.xlsx"))
BACKUP_DIR = Path(os.environ.get("PROTHEUS_BACKUP_DIR", "outputs/backups"))
API_LOG_DIR = BACKUP_DIR / "Logs API AK2_GN"

PRODUCAO_URL_BUDGET = os.environ.get("PROTHEUS_BUDGET_URL", "https://erp.example.invalid/api/budget")
PRODUCAO_URL_EXTRA = os.environ.get("PROTHEUS_EXTRA_URL", "https://erp.example.invalid/api/extra-budget")
API_USER = os.environ.get("PROTHEUS_USERNAME", "")
API_PASSWORD = os.environ.get("PROTHEUS_PASSWORD", "")

API_BATCH_SIZE = 10_000
API_TIMEOUT_SECONDS = 180
EXTRA_MAX_WORKERS = 10
EXTRA_PROGRESS_LOG_EVERY = 25

MAX_INLINE_JSON_LOG_CHARS = 12_000
MAX_INLINE_RESPONSE_LOG_CHARS = 6_000

SHEET_DADOS = "DADOS IMP"
SHEET_ITEMCC = "dItemCC"
SHEET_LOTES = "Lotes Protheus"
TABLE_LOTES = "LoteProtheus"

SHEET_EXTRA = "Orçamento Extra"
TABLE_EXTRA = "tblExtra"

MODO_BUDGET = "Budget"
MODO_FORECAST = "Forecast"

MONTH_COLUMNS = [
    ("Janeiro", 1),
    ("Fevereiro", 2),
    ("Março", 3),
    ("Abril", 4),
    ("Maio", 5),
    ("Junho", 6),
    ("Julho", 7),
    ("Agosto", 8),
    ("Setembro", 9),
    ("Outubro", 10),
    ("Novembro", 11),
    ("Dezembro", 12),
]

# novo padrão de lote: AAA001 ... ZZZ999
LOT_PREFIX_LEN = 3
LOT_NUM_LEN = 3
LOT_MAX_PER_PREFIX = 999
LOT_PREFIX_SPACE = 26 ** LOT_PREFIX_LEN
LOT_TOTAL_CAPACITY = LOT_PREFIX_SPACE * LOT_MAX_PER_PREFIX
LOT_REGEX = re.compile(r"^([A-Z]{3})(\d{3})$")


# =========================================================
# UTILIDADES GERAIS
# =========================================================
def configurar_logging(fila_gui: Queue | None = None):
    root = logging.getLogger()
    root.setLevel(logging.INFO)

    for h in list(root.handlers):
        root.removeHandler(h)

    fmt = logging.Formatter("%(asctime)s | %(levelname)-7s | %(message)s")

    ch = logging.StreamHandler(sys.stdout)
    ch.setFormatter(fmt)
    root.addHandler(ch)

    if fila_gui is not None:

        class TkQueueHandler(logging.Handler):
            def __init__(self, q: Queue):
                super().__init__(level=logging.INFO)
                self.q = q
                self.setFormatter(fmt)

            def emit(self, record):
                try:
                    self.q.put_nowait(self.format(record))
                except Exception:
                    pass

        root.addHandler(TkQueueHandler(fila_gui))

    logging.info("Logging inicializado.")
    return root


class TextWidgetLogHandler(logging.Handler):
    def __init__(self, widget: ScrolledText):
        super().__init__(level=logging.INFO)
        self.widget = widget
        self.setFormatter(logging.Formatter("%(asctime)s | %(levelname)-7s | %(message)s"))

    def emit(self, record):
        msg = self.format(record)

        def append():
            try:
                self.widget.configure(state=NORMAL)
                self.widget.insert(END, msg + ("\n" if not msg.endswith("\n") else ""))
                self.widget.see(END)
                self.widget.configure(state=DISABLED)
            except Exception:
                pass

        try:
            self.widget.after(0, append)
        except Exception:
            pass


class AtomicCounter:
    def __init__(self, initial: int = 0):
        self._value = initial
        self._lock = threading.Lock()

    def increment(self) -> int:
        with self._lock:
            self._value += 1
            return self._value

    @property
    def value(self) -> int:
        with self._lock:
            return self._value


def remover_acentos(texto: str) -> str:
    return "".join(
        c for c in unicodedata.normalize("NFD", texto)
        if unicodedata.category(c) != "Mn"
    )


def normalizar_texto(valor: Any) -> str:
    if valor is None:
        return ""
    s = str(valor).strip().lower()
    s = remover_acentos(s)
    s = re.sub(r"\s+", " ", s)
    return s


def eh_vazio(valor: Any) -> bool:
    if valor is None:
        return True
    s = str(valor).strip()
    return s == "" or s.lower() in {"nan", "none", "null"}


def to_str_limpo(valor: Any) -> str:
    if eh_vazio(valor):
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def to_float_2(valor: Any) -> float | None:
    if valor is None:
        return None
    try:
        if isinstance(valor, str):
            s = valor.strip()
            if "," in s and "." in s:
                s = s.replace(".", "").replace(",", ".")
            else:
                s = s.replace(",", ".")
            valor = s

        n = float(valor)
        if math.isnan(n) or math.isinf(n):
            return None
        return round(n, 2)
    except Exception:
        return None


def to_int(valor: Any) -> int | None:
    if eh_vazio(valor):
        return None
    try:
        s = to_str_limpo(valor).replace(".0", "")
        return int(float(s))
    except Exception:
        return None


def classificar_modo_planilha(valor_planilha: Any) -> str | None:
    s = normalizar_texto(valor_planilha)
    if "forecast" in s:
        return MODO_FORECAST
    if "budget" in s or re.search(r"\bbud\b", s):
        return MODO_BUDGET
    return None


def nome_backup(workbook_path: Path) -> str:
    ts = datetime.now().strftime("%d.%m.%y - %H.%M")
    return f"{ts} - {workbook_path.name}"


def criar_backup(workbook_path: Path) -> Path:
    if not workbook_path.exists():
        raise FileNotFoundError(f"Arquivo não encontrado: {workbook_path}")

    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    backup_path = BACKUP_DIR / nome_backup(workbook_path)
    shutil.copy2(workbook_path, backup_path)
    logging.info("Backup criado em: %s", backup_path)
    return backup_path


def validar_estrutura_zip(workbook_path: Path) -> None:
    if not workbook_path.exists():
        raise FileNotFoundError(f"Arquivo não encontrado: {workbook_path}")

    try:
        with ZipFile(workbook_path) as zf:
            nomes = set(zf.namelist())
    except BadZipFile as e:
        raise ValueError("Arquivo Excel inválido ou corrompido.") from e

    if "xl/workbook.xml" not in nomes:
        raise ValueError("Arquivo Excel inválido ou corrompido.")

    if not any(n.startswith("xl/worksheets/") for n in nomes):
        raise ValueError("Arquivo Excel inválido: nenhuma worksheet encontrada.")


def ler_header_map(ws) -> dict[str, int]:
    primeira_linha = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not primeira_linha:
        return {}

    headers: dict[str, int] = {}
    for idx, valor in enumerate(primeira_linha, start=1):
        nome = to_str_limpo(valor)
        if nome:
            headers[nome] = idx
    return headers


def validar_header_map(ws, obrigatorias: list[str]) -> dict[str, int]:
    headers = ler_header_map(ws)
    faltando = [c for c in obrigatorias if c not in headers]
    if faltando:
        raise ValueError(f"Colunas obrigatórias ausentes em '{ws.title}': {', '.join(faltando)}")
    return headers


def worksheet_tem_dados(ws) -> bool:
    linha = next(ws.iter_rows(min_row=2, max_row=2, values_only=True), None)
    return linha is not None


def obter_table_bounds(ws, table_name: str) -> tuple[int, int, int, int]:
    if table_name not in ws.tables:
        raise ValueError(f"A tabela '{table_name}' não foi encontrada na aba '{ws.title}'.")
    return range_boundaries(ws.tables[table_name].ref)


def ler_header_map_tabela(ws, table_name: str) -> dict[str, int]:
    min_col, min_row, max_col, _ = obter_table_bounds(ws, table_name)
    headers: dict[str, int] = {}
    for col in range(min_col, max_col + 1):
        nome = to_str_limpo(ws.cell(row=min_row, column=col).value)
        if nome:
            headers[nome] = col
    return headers


def validar_header_map_tabela(ws, table_name: str, obrigatorias: list[str]) -> dict[str, int]:
    headers = ler_header_map_tabela(ws, table_name)
    faltando = [c for c in obrigatorias if c not in headers]
    if faltando:
        raise ValueError(
            f"Colunas obrigatórias ausentes na tabela '{table_name}' da aba '{ws.title}': {', '.join(faltando)}"
        )
    return headers


def iterar_linhas_tabela(ws, table_name: str):
    min_col, min_row, max_col, max_row = obter_table_bounds(ws, table_name)
    for row in ws.iter_rows(
        min_row=min_row + 1,
        max_row=max_row,
        min_col=min_col,
        max_col=max_col,
        values_only=True,
    ):
        yield row


def garantir_pasta_logs_api() -> Path:
    API_LOG_DIR.mkdir(parents=True, exist_ok=True)
    return API_LOG_DIR


def slug_nome_arquivo(texto: str, limite: int = 120) -> str:
    s = remover_acentos(to_str_limpo(texto)).upper()
    s = re.sub(r"[^A-Z0-9._-]+", "_", s)
    s = s.strip("._-")
    return s[:limite] if s else "SEM_NOME"


def serializar_json_pretty(obj: Any) -> str:
    return json.dumps(obj, ensure_ascii=False, indent=2, default=str)


def salvar_texto_log_api(prefixo: str, identificador: str, extensao: str, texto: str) -> Path:
    pasta = garantir_pasta_logs_api()
    ts = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    nome = f"{ts}__{slug_nome_arquivo(prefixo)}__{slug_nome_arquivo(identificador)}.{extensao}"
    caminho = pasta / nome
    caminho.write_text(texto, encoding="utf-8")
    return caminho


def registrar_payload_api(contexto: str, identificador: str, payload: dict[str, Any]) -> Path:
    payload_txt = serializar_json_pretty(payload)
    arquivo = salvar_texto_log_api(f"{contexto}_payload", identificador, "json", payload_txt)

    preview = payload_txt
    if len(preview) > MAX_INLINE_JSON_LOG_CHARS:
        preview = (
            preview[:MAX_INLINE_JSON_LOG_CHARS]
            + f"\n...[TRUNCADO NO LOG] JSON COMPLETO SALVO EM: {arquivo}"
        )

    logging.info(
        "JSON enviado | Contexto=%s | Identificador=%s | Arquivo=%s\n%s",
        contexto,
        identificador,
        arquivo,
        preview,
    )
    return arquivo


def registrar_resposta_api(contexto: str, identificador: str, resp: requests.Response) -> Path:
    try:
        corpo = resp.text or ""
    except Exception:
        corpo = ""

    conteudo = (
        f"URL: {resp.request.url if resp.request is not None else ''}\n"
        f"METHOD: {resp.request.method if resp.request is not None else ''}\n"
        f"STATUS_CODE: {resp.status_code}\n"
        f"REASON: {resp.reason}\n"
        f"HEADERS:\n{dict(resp.headers)}\n\n"
        f"BODY:\n{corpo}"
    )

    arquivo = salvar_texto_log_api(f"{contexto}_response", identificador, "txt", conteudo)

    preview = corpo if corpo else "[sem corpo na resposta]"
    if len(preview) > MAX_INLINE_RESPONSE_LOG_CHARS:
        preview = (
            preview[:MAX_INLINE_RESPONSE_LOG_CHARS]
            + f"\n...[TRUNCADO NO LOG] RESPOSTA COMPLETA SALVA EM: {arquivo}"
        )

    logging.info(
        "Resposta da API | Contexto=%s | Identificador=%s | HTTP=%s | Arquivo=%s\n%s",
        contexto,
        identificador,
        resp.status_code,
        arquivo,
        preview,
    )
    return arquivo


# =========================================================
# VALIDAÇÕES E CARGAS - PROCESSO PRINCIPAL
# =========================================================
def validar_estrutura_rapida(workbook_path: Path) -> None:
    validar_estrutura_zip(workbook_path)

    wb = load_workbook(workbook_path, read_only=True, data_only=True, keep_links=False)
    try:
        faltando_abas = [s for s in [SHEET_DADOS, SHEET_ITEMCC, SHEET_LOTES] if s not in wb.sheetnames]
        if faltando_abas:
            raise ValueError(
                f"A planilha precisa conter obrigatoriamente as abas "
                f"'{SHEET_DADOS}', '{SHEET_ITEMCC}' e '{SHEET_LOTES}'. "
                f"Ausentes: {', '.join(faltando_abas)}"
            )

        ws_dados = wb[SHEET_DADOS]
        if not worksheet_tem_dados(ws_dados):
            raise ValueError(f"A aba '{SHEET_DADOS}' está vazia.")
        validar_header_map(
            ws_dados,
            ["Empresa", "CC", "ITEM.C", "Conta.C", "ANO_MES", "Planilha", "Valor", "Centro_Custo"],
        )

        ws_item = wb[SHEET_ITEMCC]
        if not worksheet_tem_dados(ws_item):
            raise ValueError(f"A aba '{SHEET_ITEMCC}' está vazia.")
        validar_header_map(ws_item, ["ITEM", "CC"])

        if SHEET_LOTES not in wb.sheetnames:
            raise ValueError(f"A aba '{SHEET_LOTES}' está ausente.")
    finally:
        wb.close()


@dataclass
class WorkbookCache:
    dados: list[dict[str, Any]]
    mapa_emp_item: dict[tuple[str, str], str]
    mapa_item: dict[str, str]
    centros_por_modo: dict[str, list[str]]
    empresas_por_modo: dict[str, dict[str, list[str]]]


def carregar_mapas_itemcc_de_ws(ws_item) -> tuple[dict[tuple[str, str], str], dict[str, str]]:
    headers_item = validar_header_map(ws_item, ["ITEM", "CC"])

    idx_item = {
        "ITEM": headers_item["ITEM"],
        "CC": headers_item["CC"],
        "NUM_PREFIXADO": headers_item.get("NUM_PREFIXADO"),
    }

    mapa_emp_item: dict[tuple[str, str], str] = {}
    mapa_item: dict[str, str] = {}

    for row in ws_item.iter_rows(min_row=2, values_only=True):
        if row is None or all(v is None for v in row):
            continue

        item = to_str_limpo(row[idx_item["ITEM"] - 1]) if idx_item["ITEM"] - 1 < len(row) else ""
        cc = to_str_limpo(row[idx_item["CC"] - 1]) if idx_item["CC"] - 1 < len(row) else ""

        if not item or not cc:
            continue

        num_prefixado = ""
        if idx_item["NUM_PREFIXADO"] is not None and idx_item["NUM_PREFIXADO"] - 1 < len(row):
            num_prefixado = to_str_limpo(row[idx_item["NUM_PREFIXADO"] - 1])

        empresa_prefixo = ""
        m = re.match(r"^(\d{2})\.", num_prefixado)
        if m:
            empresa_prefixo = m.group(1)

        if empresa_prefixo and (empresa_prefixo, item) not in mapa_emp_item:
            mapa_emp_item[(empresa_prefixo, item)] = cc

        if item not in mapa_item:
            mapa_item[item] = cc

    return mapa_emp_item, mapa_item


def carregar_dados_para_execucao(workbook_path: Path) -> WorkbookCache:
    wb = load_workbook(workbook_path, read_only=True, data_only=True, keep_links=False)
    try:
        ws_dados = wb[SHEET_DADOS]
        ws_item = wb[SHEET_ITEMCC]

        headers_dados = validar_header_map(
            ws_dados,
            ["Empresa", "CC", "ITEM.C", "Conta.C", "ANO_MES", "Planilha", "Valor", "Centro_Custo"],
        )

        idx_dados = {
            "Empresa": headers_dados["Empresa"],
            "CC": headers_dados["CC"],
            "ITEM.C": headers_dados["ITEM.C"],
            "Conta.C": headers_dados["Conta.C"],
            "ANO_MES": headers_dados["ANO_MES"],
            "Planilha": headers_dados["Planilha"],
            "Valor": headers_dados["Valor"],
            "Centro_Custo": headers_dados["Centro_Custo"],
        }

        dados: list[dict[str, Any]] = []
        centros_budget: set[str] = set()
        centros_forecast: set[str] = set()
        empresas_budget: dict[str, set[str]] = {}
        empresas_forecast: dict[str, set[str]] = {}

        for row in ws_dados.iter_rows(min_row=2, values_only=True):
            if row is None or all(v is None for v in row):
                continue

            registro = {
                "Empresa": row[idx_dados["Empresa"] - 1] if idx_dados["Empresa"] - 1 < len(row) else None,
                "CC": row[idx_dados["CC"] - 1] if idx_dados["CC"] - 1 < len(row) else None,
                "ITEM.C": row[idx_dados["ITEM.C"] - 1] if idx_dados["ITEM.C"] - 1 < len(row) else None,
                "Conta.C": row[idx_dados["Conta.C"] - 1] if idx_dados["Conta.C"] - 1 < len(row) else None,
                "ANO_MES": row[idx_dados["ANO_MES"] - 1] if idx_dados["ANO_MES"] - 1 < len(row) else None,
                "Planilha": row[idx_dados["Planilha"] - 1] if idx_dados["Planilha"] - 1 < len(row) else None,
                "Valor": row[idx_dados["Valor"] - 1] if idx_dados["Valor"] - 1 < len(row) else None,
                "Centro_Custo": row[idx_dados["Centro_Custo"] - 1] if idx_dados["Centro_Custo"] - 1 < len(row) else None,
            }
            dados.append(registro)

            modo = classificar_modo_planilha(registro.get("Planilha"))
            centro = to_str_limpo(registro.get("Centro_Custo"))
            empresa = to_str_limpo(registro.get("Empresa")).zfill(2)[-2:] if to_str_limpo(registro.get("Empresa")) else ""

            if not centro or modo not in (MODO_BUDGET, MODO_FORECAST):
                continue

            if modo == MODO_BUDGET:
                centros_budget.add(centro)
                empresas_budget.setdefault(empresa or "Sem empresa", set()).add(centro)
            else:
                centros_forecast.add(centro)
                empresas_forecast.setdefault(empresa or "Sem empresa", set()).add(centro)

        mapa_emp_item, mapa_item = carregar_mapas_itemcc_de_ws(ws_item)
    finally:
        wb.close()

    return WorkbookCache(
        dados=dados,
        mapa_emp_item=mapa_emp_item,
        mapa_item=mapa_item,
        centros_por_modo={
            MODO_BUDGET: sorted(centros_budget),
            MODO_FORECAST: sorted(centros_forecast),
        },
        empresas_por_modo={
            MODO_BUDGET: {k: sorted(v) for k, v in sorted(empresas_budget.items())},
            MODO_FORECAST: {k: sorted(v) for k, v in sorted(empresas_forecast.items())},
        },
    )


def extrair_company_code(centro_custo: Any, empresa: Any) -> str | None:
    s = to_str_limpo(centro_custo)
    m = re.search(r"(\d{2})", s)
    if m:
        return f"{m.group(1)}01"

    emp = to_str_limpo(empresa)
    m = re.search(r"(\d{2})", emp)
    if m:
        return f"{m.group(1)}01"

    return None


def preencher_cc_se_vazio(
    row: dict[str, Any],
    mapa_emp_item: dict[tuple[str, str], str],
    mapa_item: dict[str, str],
) -> str:
    cc_atual = to_str_limpo(row.get("CC"))
    if cc_atual:
        return cc_atual

    item = to_str_limpo(row.get("ITEM.C"))
    empresa = to_str_limpo(row.get("Empresa"))

    if item and empresa:
        chave = (empresa.zfill(2)[-2:], item)
        if chave in mapa_emp_item:
            return mapa_emp_item[chave]

    if item and item in mapa_item:
        return mapa_item[item]

    return cc_atual


@dataclass
class RegistroApi:
    company_code: str
    centro_de_custo: str
    item_contabil: str
    conta: str
    competencia: str
    valor: float


def preparar_registros_api(
    dados: list[dict[str, Any]],
    mapa_emp_item: dict[tuple[str, str], str],
    mapa_item: dict[str, str],
    modo: str,
    centros_selecionados: list[str],
) -> list[RegistroApi]:
    selecionados_norm = {normalizar_texto(v) for v in centros_selecionados if to_str_limpo(v)}

    registros: list[RegistroApi] = []
    total_modo = 0
    total_filtro = 0
    total_cc_preenchido = 0
    total_descartados = 0

    for row in dados:
        if classificar_modo_planilha(row.get("Planilha")) != modo:
            continue
        total_modo += 1

        centro_custo_original = to_str_limpo(row.get("Centro_Custo"))
        if selecionados_norm and normalizar_texto(centro_custo_original) not in selecionados_norm:
            continue
        total_filtro += 1

        cc_original = to_str_limpo(row.get("CC"))
        cc_final = preencher_cc_se_vazio(row, mapa_emp_item, mapa_item)
        if not cc_original and cc_final:
            total_cc_preenchido += 1

        company_code = extrair_company_code(centro_custo_original, row.get("Empresa"))
        competencia = to_str_limpo(row.get("ANO_MES"))
        conta = to_str_limpo(row.get("Conta.C"))
        item_contabil = to_str_limpo(row.get("ITEM.C"))
        valor = to_float_2(row.get("Valor"))

        if not company_code or not competencia or not cc_final or not conta or valor is None:
            total_descartados += 1
            continue

        registros.append(
            RegistroApi(
                company_code=company_code,
                centro_de_custo=cc_final,
                item_contabil=item_contabil,
                conta=conta,
                competencia=competencia,
                valor=valor,
            )
        )

    logging.info("Registros lidos para o modo %s: %d", modo, total_modo)
    logging.info("Registros após filtro de Centro_Custo: %d", total_filtro)
    logging.info("CC preenchidos via dItemCC: %d", total_cc_preenchido)
    logging.info("Registros descartados por campos obrigatórios: %d", total_descartados)
    logging.info("Registros prontos para API: %d", len(registros))
    return registros


def agrupar_por_company(registros: list[RegistroApi]) -> dict[str, list[RegistroApi]]:
    grupos: dict[str, list[RegistroApi]] = {}
    for reg in registros:
        grupos.setdefault(reg.company_code, []).append(reg)
    return grupos


# =========================================================
# LÓGICA DE LOTES - PROCESSO PRINCIPAL
# =========================================================
def prefixo_para_indice(prefixo: str) -> int:
    prefixo = prefixo.upper().strip()
    if len(prefixo) != 3 or not prefixo.isalpha():
        raise ValueError(f"Prefixo inválido: {prefixo}")

    idx = 0
    for ch in prefixo:
        idx = idx * 26 + (ord(ch) - ord("A"))
    return idx


def indice_para_prefixo(indice: int) -> str:
    if indice < 0 or indice >= LOT_PREFIX_SPACE:
        raise ValueError(f"Índice de prefixo fora do intervalo: {indice}")

    chars = []
    n = indice
    for _ in range(LOT_PREFIX_LEN):
        chars.append(chr(ord("A") + (n % 26)))
        n //= 26
    return "".join(reversed(chars))


def lote_para_indice_global(lote: str) -> int | None:
    lote = to_str_limpo(lote).upper()
    m = LOT_REGEX.match(lote)
    if not m:
        return None

    prefixo = m.group(1)
    numero = int(m.group(2))

    if numero < 1 or numero > LOT_MAX_PER_PREFIX:
        return None

    prefix_idx = prefixo_para_indice(prefixo)
    return prefix_idx * LOT_MAX_PER_PREFIX + (numero - 1)


def indice_global_para_lote(indice_global: int) -> str:
    if indice_global < 0 or indice_global >= LOT_TOTAL_CAPACITY:
        raise ValueError("Capacidade máxima de lotes atingida.")

    prefix_idx = indice_global // LOT_MAX_PER_PREFIX
    numero = (indice_global % LOT_MAX_PER_PREFIX) + 1

    prefixo = indice_para_prefixo(prefix_idx)
    return f"{prefixo}{numero:03d}"


@dataclass
class LoteSeqInfoFast:
    proximo_indice_global: int


def obter_proximos_lotes_fast(workbook_path: Path) -> LoteSeqInfoFast:
    logging.info("Lendo sequência atual de lotes de forma otimizada...")

    wb = load_workbook(workbook_path, read_only=True, data_only=True, keep_links=False)
    try:
        if SHEET_LOTES not in wb.sheetnames:
            raise ValueError(f"A aba '{SHEET_LOTES}' não foi encontrada.")

        ws = wb[SHEET_LOTES]

        maior_indice_global = -1
        header_row = None
        col_lote = None

        for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=20, values_only=True), start=1):
            valores = [to_str_limpo(v) for v in row]
            if "Lote" in valores:
                header_row = r_idx
                col_lote = valores.index("Lote") + 1
                break

        if header_row is None or col_lote is None:
            raise ValueError(
                f"Não foi possível localizar a coluna 'Lote' na aba '{SHEET_LOTES}'."
            )

        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            if row is None:
                continue
            if col_lote - 1 >= len(row):
                continue

            valor = to_str_limpo(row[col_lote - 1]).upper()
            if not valor:
                continue

            idx_global = lote_para_indice_global(valor)
            if idx_global is not None:
                maior_indice_global = max(maior_indice_global, idx_global)

        proximo = maior_indice_global + 1
        if proximo >= LOT_TOTAL_CAPACITY:
            raise RuntimeError("Capacidade máxima de lotes atingida (ZZZ999).")

        logging.info(
            "Maior lote encontrado no novo padrão: %s | Próximo lote: %s",
            indice_global_para_lote(maior_indice_global) if maior_indice_global >= 0 else "[nenhum]",
            indice_global_para_lote(proximo),
        )

        return LoteSeqInfoFast(proximo_indice_global=proximo)
    finally:
        wb.close()


def append_lotes_via_excel_com(workbook_path: Path, lotes: list[tuple[str, int]]) -> None:
    if not lotes:
        return

    logging.info("Abrindo Excel via COM para registrar %d lote(s)...", len(lotes))

    pythoncom.CoInitialize()
    excel = win32.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    excel.ScreenUpdating = False
    excel.EnableEvents = False
    excel.AskToUpdateLinks = False

    wb = None
    try:
        wb = excel.Workbooks.Open(str(workbook_path), UpdateLinks=0, ReadOnly=False)
        ws = wb.Worksheets(SHEET_LOTES)

        try:
            tbl = ws.ListObjects(TABLE_LOTES)
        except Exception as e:
            raise RuntimeError(
                f"Não foi possível localizar a tabela '{TABLE_LOTES}' na aba '{SHEET_LOTES}'."
            ) from e

        header_range = tbl.HeaderRowRange
        headers_map: dict[str, int] = {}
        for col in range(1, header_range.Columns.Count + 1):
            val = header_range.Cells(1, col).Value
            nome = to_str_limpo(val)
            if nome:
                headers_map[nome] = col

        if "Lote" not in headers_map or "Qtd Registros" not in headers_map:
            raise RuntimeError(
                f"A tabela '{TABLE_LOTES}' deve conter as colunas 'Lote' e 'Qtd Registros'."
            )

        col_lote = headers_map["Lote"]
        col_qtd = headers_map["Qtd Registros"]

        for lote, qtd in lotes:
            row_obj = tbl.ListRows.Add()
            rng = row_obj.Range
            rng.Cells(1, col_lote).Value = lote
            rng.Cells(1, col_qtd).Value = qtd

        wb.Save()
        logging.info("Lotes registrados com sucesso via Excel COM.")
    finally:
        if wb is not None:
            try:
                wb.Close(SaveChanges=False)
            except Exception:
                pass
        try:
            excel.Quit()
        except Exception:
            pass
        pythoncom.CoUninitialize()


def limpar_tabela_via_excel_com(workbook_path: Path, sheet_name: str, table_name: str) -> None:
    logging.info("Abrindo Excel via COM para limpar a tabela '%s'...", table_name)

    pythoncom.CoInitialize()
    excel = win32.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    excel.ScreenUpdating = False
    excel.EnableEvents = False
    excel.AskToUpdateLinks = False

    wb = None
    try:
        wb = excel.Workbooks.Open(str(workbook_path), UpdateLinks=0, ReadOnly=False)
        ws = wb.Worksheets(sheet_name)

        try:
            tbl = ws.ListObjects(table_name)
        except Exception as e:
            raise RuntimeError(
                f"Não foi possível localizar a tabela '{table_name}' na aba '{sheet_name}'."
            ) from e

        while tbl.ListRows.Count > 0:
            tbl.ListRows(tbl.ListRows.Count).Delete()

        wb.Save()
        logging.info("Tabela '%s' limpa com sucesso.", table_name)
    finally:
        if wb is not None:
            try:
                wb.Close(SaveChanges=False)
            except Exception:
                pass
        try:
            excel.Quit()
        except Exception:
            pass
        pythoncom.CoUninitialize()


def montar_nome_lote(indice_global: int) -> str:
    return indice_global_para_lote(indice_global)


def enviar_lotes_e_registrar(
    workbook_path: Path,
    registros: list[RegistroApi],
    modo: str,
) -> tuple[int, int]:
    if not registros:
        logging.info("Nenhum registro para envio.")
        return 0, 0

    grupos = agrupar_por_company(registros)

    seq_fast = obter_proximos_lotes_fast(workbook_path)
    indice_global_atual = seq_fast.proximo_indice_global

    logging.info("Iniciando envio HTTP para Produção...")
    logging.info("Pasta de logs da API: %s", garantir_pasta_logs_api())

    session = requests.Session()
    session.auth = (API_USER, API_PASSWORD)
    session.headers.update({
        "Content-Type": "application/json",
        "Accept": "application/json",
    })

    total_lotes_ok = 0
    total_registros_ok = 0
    lotes_para_registrar: list[tuple[str, int]] = []

    for company_code in sorted(grupos.keys()):
        regs = grupos[company_code]

        for ini in range(0, len(regs), API_BATCH_SIZE):
            chunk = regs[ini:ini + API_BATCH_SIZE]
            nome_lote = montar_nome_lote(indice_global_atual)
            indice_global_atual += 1

            payload = {
                "CompanyCode": company_code,
                "IdProcessamento": nome_lote,
                "Registros": [
                    {
                        "Competencia": reg.competencia,
                        "CentroDeCusto": reg.centro_de_custo,
                        "ItemContabil": reg.item_contabil,
                        "Conta": reg.conta,
                        "Valor": reg.valor,
                        "Indexador": idx,
                    }
                    for idx, reg in enumerate(chunk, start=1)
                ],
            }

            logging.info(
                "Enviando lote %s | Modo=%s | Ambiente=Produção | CompanyCode=%s | Qtd=%d",
                nome_lote,
                modo,
                company_code,
                len(chunk),
            )

            payload_file = registrar_payload_api("budget", nome_lote, payload)

            try:
                resp = session.post(
                    PRODUCAO_URL_BUDGET,
                    json=payload,
                    timeout=API_TIMEOUT_SECONDS,
                    verify=True,
                )
            except requests.RequestException as e:
                raise RuntimeError(
                    f"Erro ao enviar lote {nome_lote}: {e}\nPayload salvo em: {payload_file}"
                ) from e

            resposta_file = registrar_resposta_api("budget", nome_lote, resp)

            if resp.status_code != 200:
                raise RuntimeError(
                    f"Falha ao enviar lote {nome_lote}. "
                    f"HTTP {resp.status_code}. "
                    f"Payload: {payload_file} | Resposta: {resposta_file}"
                )

            lotes_para_registrar.append((nome_lote, len(chunk)))
            total_lotes_ok += 1
            total_registros_ok += len(chunk)
            logging.info("Lote %s enviado com sucesso.", nome_lote)

    logging.info("Todos os lotes enviados. Registrando lotes na planilha via Excel COM...")
    append_lotes_via_excel_com(workbook_path, lotes_para_registrar)

    return total_lotes_ok, total_registros_ok


# =========================================================
# EXTRA ORÇAMENTO - SOMENTE IMPORTAÇÃO DE PLANILHA
# =========================================================
def validar_estrutura_extra(workbook_path: Path) -> None:
    validar_estrutura_zip(workbook_path)

    wb = load_workbook(workbook_path, read_only=False, data_only=True, keep_links=False)
    try:
        faltando_abas = [s for s in [SHEET_EXTRA, SHEET_ITEMCC] if s not in wb.sheetnames]
        if faltando_abas:
            raise ValueError(
                f"A planilha de Extra Orçamento precisa conter as abas "
                f"'{SHEET_EXTRA}' e '{SHEET_ITEMCC}'. Ausentes: {', '.join(faltando_abas)}"
            )

        ws_extra = wb[SHEET_EXTRA]
        ws_item = wb[SHEET_ITEMCC]

        if TABLE_EXTRA not in ws_extra.tables:
            raise ValueError(f"A aba '{SHEET_EXTRA}' precisa conter a tabela '{TABLE_EXTRA}'.")

        validar_header_map_tabela(
            ws_extra,
            TABLE_EXTRA,
            ["EMP", "Operação", "C.C", "Item.C", "Conta", "Código Conta", "Ano"],
        )
        validar_header_map(ws_item, ["ITEM", "CC"])
    finally:
        wb.close()


def extrair_company_code_por_emp(emp: Any) -> str | None:
    s = to_str_limpo(emp)
    m = re.search(r"(\d{2})", s.zfill(2))
    if m:
        return f"{m.group(1)}01"
    return None


def resolver_cc_por_cc_ou_item(
    empresa: Any,
    cc: Any,
    item: Any,
    mapa_emp_item: dict[tuple[str, str], str],
    mapa_item: dict[str, str],
) -> str:
    cc_atual = to_str_limpo(cc)
    if cc_atual:
        return cc_atual

    item_limpo = to_str_limpo(item)
    empresa_limpa = to_str_limpo(empresa).zfill(2)[-2:]

    if item_limpo and empresa_limpa:
        chave = (empresa_limpa, item_limpo)
        if chave in mapa_emp_item:
            return mapa_emp_item[chave]

    if item_limpo and item_limpo in mapa_item:
        return mapa_item[item_limpo]

    return ""


@dataclass
class ExtraRegistroApi:
    company_code: str
    competencia: str
    centro_de_custo: str
    conta: str
    valor: float
    linha_origem: int


def preparar_registros_extra_de_planilha(workbook_path: Path) -> list[ExtraRegistroApi]:
    wb = load_workbook(workbook_path, read_only=False, data_only=True, keep_links=False)
    try:
        if SHEET_EXTRA not in wb.sheetnames:
            raise ValueError(f"A aba '{SHEET_EXTRA}' não foi encontrada.")
        if SHEET_ITEMCC not in wb.sheetnames:
            raise ValueError(f"A aba '{SHEET_ITEMCC}' não foi encontrada.")

        ws_extra = wb[SHEET_EXTRA]
        ws_item = wb[SHEET_ITEMCC]

        headers_extra = validar_header_map_tabela(
            ws_extra,
            TABLE_EXTRA,
            ["EMP", "Operação", "C.C", "Item.C", "Conta", "Código Conta", "Ano"],
        )
        _ = headers_extra

        mapa_emp_item, mapa_item = carregar_mapas_itemcc_de_ws(ws_item)

        min_col, min_row, max_col, _ = obter_table_bounds(ws_extra, TABLE_EXTRA)
        ordered_headers = []
        for col in range(min_col, max_col + 1):
            ordered_headers.append(to_str_limpo(ws_extra.cell(row=min_row, column=col).value))

        registros: list[ExtraRegistroApi] = []
        total_linhas = 0
        total_competencias = 0
        total_descartados = 0
        total_cc_via_item = 0

        for idx_linha, row in enumerate(iterar_linhas_tabela(ws_extra, TABLE_EXTRA), start=min_row + 1):
            if row is None or all(eh_vazio(v) for v in row):
                continue

            total_linhas += 1
            row_map = {ordered_headers[idx]: row[idx] if idx < len(row) else None for idx in range(len(ordered_headers))}

            emp = to_str_limpo(row_map.get("EMP"))
            cc = to_str_limpo(row_map.get("C.C"))
            item = to_str_limpo(row_map.get("Item.C"))
            conta_codigo = to_str_limpo(row_map.get("Código Conta"))
            ano = to_int(row_map.get("Ano"))

            cc_final = resolver_cc_por_cc_ou_item(emp, cc, item, mapa_emp_item, mapa_item)
            if not cc and cc_final:
                total_cc_via_item += 1

            company_code = extrair_company_code_por_emp(emp)

            if not company_code or not cc_final or not conta_codigo or not ano:
                total_descartados += 1
                logging.warning(
                    "Linha de Extra descartada por campos obrigatórios inválidos | Linha=%s | EMP=%s | C.C=%s | Item.C=%s | Código Conta=%s | Ano=%s",
                    idx_linha,
                    emp,
                    cc,
                    item,
                    conta_codigo,
                    row_map.get("Ano"),
                )
                continue

            teve_valor = False
            for month_name, month_num in MONTH_COLUMNS:
                valor = to_float_2(row_map.get(month_name))
                if valor is None or valor <= 0:
                    continue

                teve_valor = True
                total_competencias += 1
                registros.append(
                    ExtraRegistroApi(
                        company_code=company_code,
                        competencia=f"{ano:04d}{month_num:02d}",
                        centro_de_custo=cc_final,
                        conta=conta_codigo,
                        valor=valor,
                        linha_origem=idx_linha,
                    )
                )

            if not teve_valor:
                logging.info(
                    "Linha de Extra sem valores positivos ignorada | Linha=%s | EMP=%s | C.C=%s | Item.C=%s | Conta=%s",
                    idx_linha,
                    emp,
                    cc_final,
                    item,
                    conta_codigo,
                )

        logging.info("Linhas lidas na tabela %s: %d", TABLE_EXTRA, total_linhas)
        logging.info("CC resolvidos via dItemCC na planilha de Extra: %d", total_cc_via_item)
        logging.info("Linhas descartadas na planilha de Extra: %d", total_descartados)
        logging.info("Registros/competências prontos para API de Extra: %d", len(registros))
        logging.info("Competências positivas identificadas na planilha: %d", total_competencias)

        return registros
    finally:
        wb.close()


def _enviar_um_extra(
    indice: int,
    total: int,
    reg: ExtraRegistroApi,
    progress_counter: AtomicCounter,
) -> int:
    payload = {
        "CompanyCode": reg.company_code,
        "Competencia": reg.competencia,
        "CentroDeCusto": reg.centro_de_custo,
        "Conta": reg.conta,
        "Valor": reg.valor,
    }

    identificador = f"extra_{indice:05d}_linha_{reg.linha_origem}_{reg.company_code}_{reg.competencia}"
    payload_file = registrar_payload_api("extra", identificador, payload)

    session = requests.Session()
    session.auth = (API_USER, API_PASSWORD)
    session.headers.update({
        "Content-Type": "application/json; charset=utf-8",
        "Accept": "application/json",
    })

    try:
        resp = session.post(
            PRODUCAO_URL_EXTRA,
            json=payload,
            timeout=API_TIMEOUT_SECONDS,
            verify=True,
        )
    except requests.RequestException as e:
        raise RuntimeError(
            f"Erro ao enviar Extra {indice}/{total} | Linha origem={reg.linha_origem} | "
            f"Competencia={reg.competencia} | CC={reg.centro_de_custo}: {e} | "
            f"Payload salvo em: {payload_file}"
        ) from e
    finally:
        session.close()

    resposta_file = registrar_resposta_api("extra", identificador, resp)

    if resp.status_code not in (200, 201, 204):
        raise RuntimeError(
            f"Falha ao enviar Extra {indice}/{total} | Linha origem={reg.linha_origem} | "
            f"HTTP {resp.status_code}. Payload: {payload_file} | Resposta: {resposta_file}"
        )

    concluidos = progress_counter.increment()
    if concluidos == 1 or concluidos == total or concluidos % EXTRA_PROGRESS_LOG_EVERY == 0:
        logging.info(
            "Progresso Extra: %d/%d enviados com sucesso.",
            concluidos,
            total,
        )

    return 1


def enviar_extras_paralelo(registros: list[ExtraRegistroApi], max_workers: int = EXTRA_MAX_WORKERS) -> int:
    if not registros:
        logging.info("Nenhum extra elegível para envio.")
        return 0

    total = len(registros)
    workers = max(1, min(max_workers, total))
    logging.info(
        "Iniciando envio de Extra Orçamento para Produção com %d thread(s). Total de registros: %d",
        workers,
        total,
    )
    logging.info("Pasta de logs da API: %s", garantir_pasta_logs_api())

    progress_counter = AtomicCounter(0)
    total_ok = 0
    falhas: list[str] = []

    with ThreadPoolExecutor(max_workers=workers, thread_name_prefix="extra_api") as executor:
        future_map = {
            executor.submit(_enviar_um_extra, idx, total, reg, progress_counter): (idx, reg)
            for idx, reg in enumerate(registros, start=1)
        }

        for future in as_completed(future_map):
            idx, reg = future_map[future]
            try:
                total_ok += future.result()
            except Exception as exc:
                msg = (
                    f"Falha no Extra {idx}/{total} | Linha origem={reg.linha_origem} | "
                    f"Competencia={reg.competencia} | CC={reg.centro_de_custo} | Conta={reg.conta} | Erro={exc}"
                )
                falhas.append(msg)
                logging.error(msg)

    if falhas:
        raise RuntimeError(
            "Ocorreram falhas no envio do Extra Orçamento. "
            f"Sucessos: {total_ok} | Falhas: {len(falhas)}\n\n" + "\n".join(falhas[:20])
        )

    logging.info("Envio de Extra Orçamento concluído com sucesso. Total enviado: %d", total_ok)
    return total_ok


# =========================================================
# DIÁLOGO DE SELEÇÃO DE CENTROS
# =========================================================
class SelecionarCentrosDialog(Toplevel):
    def __init__(
        self,
        master,
        valores: list[str],
        empresas_map: dict[str, list[str]],
        selecionados_iniciais: list[str] | None = None,
    ):
        super().__init__(master)
        self.title("Selecionar Centro(s) de Custo")
        self.geometry("1140x640")
        self.minsize(980, 560)
        self.transient(master)
        self.grab_set()

        self.resultado: list[str] | None = None
        self.valores = sorted({v for v in valores if to_str_limpo(v)})
        self.empresas_map = empresas_map
        self.selecionados = sorted({v for v in (selecionados_iniciais or []) if v in self.valores})

        self.busca_var = StringVar(value="")
        self.busca_var.trace_add("write", lambda *_: self._atualizar_disponiveis())

        self.empresa_var = StringVar(value="")

        top = ttk.Frame(self)
        top.pack(fill="x", padx=12, pady=(12, 8))

        ttk.Label(top, text="PESQUISAR:").pack(side="left")
        self.entry_busca = ttk.Entry(top, textvariable=self.busca_var)
        self.entry_busca.pack(side="left", fill="x", expand=True, padx=(6, 12))

        ttk.Label(top, text="EMP:").pack(side="left")
        empresas = [""] + list(self.empresas_map.keys())
        self.cbo_empresa = ttk.Combobox(
            top,
            textvariable=self.empresa_var,
            values=empresas,
            width=18,
            state="readonly",
        )
        self.cbo_empresa.pack(side="left")
        ttk.Button(top, text="ADD EMPRESA", command=self._adicionar_empresa).pack(side="left", padx=(8, 0))

        help_fr = ttk.Frame(self)
        help_fr.pack(fill="x", padx=12, pady=(0, 6))
        ttk.Label(
            help_fr,
            text="A busca aceita vários termos. Ex.: 'ambev 01 2026' retorna itens que contenham todos os termos.",
        ).pack(anchor="w")

        body = ttk.Frame(self)
        body.pack(fill=BOTH, expand=True, padx=12, pady=(0, 8))

        left = ttk.Frame(body)
        left.pack(side="left", fill=BOTH, expand=True)
        ttk.Label(left, text="DISPONÍVEIS").pack(anchor="w")
        self.lst_disp = Listbox(left, selectmode=EXTENDED)
        self.lst_disp.pack(fill=BOTH, expand=True, pady=(4, 0))
        self.lst_disp.bind("<Double-Button-1>", lambda e: self._adicionar_selecionados())

        mid = ttk.Frame(body)
        mid.pack(side="left", fill="y", padx=10)
        ttk.Button(mid, text="ADICIONAR >", width=22, command=self._adicionar_selecionados).pack(pady=(50, 6))
        ttk.Button(mid, text="ADICIONAR TODOS", width=22, command=self._adicionar_todos_filtrados).pack(pady=6)
        ttk.Button(mid, text="< REMOVER", width=22, command=self._remover_selecionados).pack(pady=6)
        ttk.Button(mid, text="LIMPAR", width=22, command=self._limpar_todos).pack(pady=6)

        right = ttk.Frame(body)
        right.pack(side="left", fill=BOTH, expand=True)
        ttk.Label(right, text="Selecionados").pack(anchor="w")
        self.lst_sel = Listbox(right, selectmode=EXTENDED)
        self.lst_sel.pack(fill=BOTH, expand=True, pady=(4, 0))
        self.lst_sel.bind("<Double-Button-1>", lambda e: self._remover_selecionados())

        bottom = ttk.Frame(self)
        bottom.pack(fill="x", padx=12, pady=(0, 12))
        self.lbl_info = ttk.Label(bottom, text="")
        self.lbl_info.pack(side="left")
        ttk.Button(bottom, text="CANCELAR", width=12, command=self._cancelar).pack(side="right")
        ttk.Button(bottom, text="CONFIRMAR", width=12, command=self._confirmar).pack(side="right", padx=(0, 8))

        self._atualizar_disponiveis()
        self._atualizar_selecionados()
        self.entry_busca.focus_set()
        self.protocol("WM_DELETE_WINDOW", self._cancelar)

    def _valores_filtrados(self) -> list[str]:
        disponiveis = [v for v in self.valores if v not in self.selecionados]
        termo = normalizar_texto(self.busca_var.get())
        if not termo:
            return disponiveis

        tokens = [t for t in termo.split(" ") if t]
        filtrados = []
        for valor in disponiveis:
            base = normalizar_texto(valor)
            if all(token in base for token in tokens):
                filtrados.append(valor)
        return filtrados

    def _atualizar_disponiveis(self):
        self.lst_disp.delete(0, END)
        for v in self._valores_filtrados():
            self.lst_disp.insert(END, v)
        self._atualizar_info()

    def _atualizar_selecionados(self):
        self.lst_sel.delete(0, END)
        for v in self.selecionados:
            self.lst_sel.insert(END, v)
        self._atualizar_info()

    def _atualizar_info(self):
        self.lbl_info.configure(
            text=f"Disponíveis filtrados: {self.lst_disp.size()} | Selecionados: {len(self.selecionados)}"
        )

    def _adicionar_selecionados(self):
        itens = [self.lst_disp.get(i) for i in self.lst_disp.curselection()]
        for item in itens:
            if item not in self.selecionados:
                self.selecionados.append(item)
        self.selecionados.sort()
        self._atualizar_disponiveis()
        self._atualizar_selecionados()

    def _adicionar_todos_filtrados(self):
        for item in self._valores_filtrados():
            if item not in self.selecionados:
                self.selecionados.append(item)
        self.selecionados.sort()
        self._atualizar_disponiveis()
        self._atualizar_selecionados()

    def _adicionar_empresa(self):
        empresa = self.empresa_var.get().strip()
        if not empresa:
            return

        for item in self.empresas_map.get(empresa, []):
            if item not in self.selecionados:
                self.selecionados.append(item)

        self.selecionados.sort()
        self._atualizar_disponiveis()
        self._atualizar_selecionados()

    def _remover_selecionados(self):
        itens = {self.lst_sel.get(i) for i in self.lst_sel.curselection()}
        self.selecionados = [v for v in self.selecionados if v not in itens]
        self._atualizar_disponiveis()
        self._atualizar_selecionados()

    def _limpar_todos(self):
        self.selecionados = []
        self._atualizar_disponiveis()
        self._atualizar_selecionados()

    def _confirmar(self):
        self.resultado = list(self.selecionados)
        self.destroy()

    def _cancelar(self):
        self.resultado = None
        self.destroy()


# =========================================================
# TELA EXTRA ORÇAMENTO - SOMENTE IMPORTAÇÃO
# =========================================================
class ExtraOrcamentoDialog(Toplevel):
    def __init__(self, master):
        super().__init__(master)
        self.title("Lançamento de Extra Orçamento")
        self.geometry("980x620")
        self.minsize(920, 580)
        self.transient(master)
        self.grab_set()

        self._worker: threading.Thread | None = None
        self._log_handler = None

        self.arquivo_var = StringVar(value=str(DEFAULT_EXTRA_WORKBOOK))
        self.status_var = StringVar(value="Selecione a planilha modelo e clique em importar para enviar os extras.")

        self._construir_ui()
        self._instalar_log_handler()
        self.protocol("WM_DELETE_WINDOW", self._fechar)

    def _instalar_log_handler(self):
        self._log_handler = TextWidgetLogHandler(self.logbox)
        logging.getLogger().addHandler(self._log_handler)

    def _remover_log_handler(self):
        if self._log_handler is not None:
            try:
                logging.getLogger().removeHandler(self._log_handler)
            except Exception:
                pass
            self._log_handler = None

    def _construir_ui(self):
        top = ttk.Frame(self)
        top.pack(fill="x", padx=14, pady=(10, 6))
        ttk.Label(top, text="LANÇAMENTO DE EXTRA ORÇAMENTO", style="Title.TLabel").pack(anchor="w")

        arq_fr = ttk.LabelFrame(self, text="Importação de planilha")
        arq_fr.pack(fill="x", padx=14, pady=(0, 8))

        row_arq = ttk.Frame(arq_fr)
        row_arq.pack(fill="x", padx=8, pady=8)
        ttk.Entry(row_arq, textvariable=self.arquivo_var).pack(side="left", fill="x", expand=True)
        ttk.Button(row_arq, text="PROCURAR", width=14, command=self._escolher_planilha).pack(side="left", padx=(6, 0))
        self.btn_importar_planilha = ttk.Button(
            row_arq,
            text="IMPORTAR E ENVIAR PLANILHA",
            width=28,
            command=self._importar_planilha,
        )
        self.btn_importar_planilha.pack(side="left", padx=(6, 0))

        ttk.Label(
            arq_fr,
            text=(
                "A planilha deve conter a aba 'Orçamento Extra', a tabela 'tblExtra' e a aba 'dItemCC'. "
                "Após sucesso total do envio, a tblExtra será limpa. O envio é feito em paralelo com até 10 threads."
            ),
        ).pack(anchor="w", padx=8, pady=(0, 8))

        status_fr = ttk.Frame(self)
        status_fr.pack(fill="x", padx=14, pady=(0, 4))
        ttk.Label(status_fr, textvariable=self.status_var).pack(anchor="w")

        btn_fr = ttk.Frame(self)
        btn_fr.pack(fill="x", padx=14, pady=(0, 8))
        ttk.Button(btn_fr, text="FECHAR", width=12, command=self._fechar).pack(side="right")

        log_fr = ttk.LabelFrame(self, text="Log do Extra Orçamento")
        log_fr.pack(fill=BOTH, expand=True, padx=14, pady=(0, 12))
        self.logbox = ScrolledText(log_fr, height=16, font=("Consolas", 10))
        self.logbox.pack(fill=BOTH, expand=True, padx=6, pady=6)
        self.logbox.configure(state=DISABLED)

    def _set_busy(self, busy: bool):
        state = DISABLED if busy else NORMAL
        self.btn_importar_planilha.configure(state=state)

    def _escolher_planilha(self):
        inicial = str(Path(self.arquivo_var.get()).parent) if self.arquivo_var.get().strip() else str(DEFAULT_EXTRA_WORKBOOK.parent)
        arq = filedialog.askopenfilename(
            initialdir=inicial or None,
            title="Escolha a planilha de Extra Orçamento",
            filetypes=[("Arquivos Excel", "*.xlsx")],
        )
        if arq:
            self.arquivo_var.set(arq)

    def _caminho_planilha(self) -> Path:
        caminho = self.arquivo_var.get().strip()
        if not caminho:
            raise ValueError("Selecione a planilha de Extra Orçamento.")
        p = Path(caminho)
        if not p.exists():
            raise FileNotFoundError(f"Arquivo não encontrado: {p}")
        if p.suffix.lower() != ".xlsx":
            raise ValueError("Selecione um arquivo .xlsx.")
        return p

    def _importar_planilha(self):
        if self._worker and self._worker.is_alive():
            return

        try:
            workbook_path = self._caminho_planilha()
        except Exception as e:
            messagebox.showerror("Erro", str(e), parent=self)
            return

        self.status_var.set("Validando, criando backup, preparando e enviando Extra Orçamento da planilha...")
        self._set_busy(True)

        def run():
            try:
                logging.info("Iniciando processo de importação de Extra Orçamento: %s", workbook_path)
                validar_estrutura_extra(workbook_path)
                criar_backup(workbook_path)
                registros = preparar_registros_extra_de_planilha(workbook_path)

                if not registros:
                    raise ValueError("Nenhum lançamento de Extra elegível foi encontrado na planilha.")

                qtd = enviar_extras_paralelo(registros, max_workers=EXTRA_MAX_WORKERS)
                limpar_tabela_via_excel_com(workbook_path, SHEET_EXTRA, TABLE_EXTRA)

                def done():
                    self.status_var.set("Importação e envio do Extra Orçamento concluídos com sucesso.")
                    self._set_busy(False)
                    messagebox.showinfo(
                        "Extra Orçamento",
                        f"Planilha importada com sucesso.\nRegistros enviados: {qtd}\nA tabela {TABLE_EXTRA} foi limpa.",
                        parent=self,
                    )

                self.after(0, done)
            except Exception as exc:
                logging.exception("Falha no processo de Extra Orçamento por planilha: %s", exc)

                def fail():
                    self.status_var.set("Falha no envio do Extra Orçamento por planilha.")
                    self._set_busy(False)
                    messagebox.showerror("Erro no Extra Orçamento", str(exc), parent=self)

                self.after(0, fail)

        self._worker = threading.Thread(target=run, daemon=True)
        self._worker.start()

    def _fechar(self):
        if self._worker and self._worker.is_alive():
            messagebox.showwarning("Processando", "Existe um processamento em andamento.", parent=self)
            return
        self._remover_log_handler()
        self.destroy()


# =========================================================
# APLICAÇÃO PRINCIPAL
# =========================================================
class App(Tk):
    def __init__(self):
        super().__init__()
        self.title(APP_TITLE)
        self.geometry("1180x780")
        self.minsize(1000, 720)

        style = ttk.Style(self)
        try:
            style.theme_use("clam")
        except Exception:
            pass
        style.configure("Title.TLabel", font=("Segoe UI", 14, "bold"))
        style.configure("TButton", font=("Segoe UI", 10))
        style.configure("TLabel", font=("Segoe UI", 10))

        self.log_queue: Queue = Queue()
        configurar_logging(self.log_queue)

        self.arquivo_var = StringVar(value=str(DEFAULT_WORKBOOK))
        self.modo_var = StringVar(value=MODO_FORECAST)
        self.status_var = StringVar(value="Aguardando carregamento da planilha padrão...")

        self._worker: threading.Thread | None = None
        self._cache: WorkbookCache | None = None
        self._arquivo_carregado: Path | None = None
        self._centros_selecionados: list[str] = []

        self._construir_ui()
        self.after(120, self._poll_log)
        self.after(250, self._carregar_padrao_automaticamente)

    def _construir_ui(self):
        top = ttk.Frame(self)
        top.pack(fill="x", padx=14, pady=(10, 4))
        ttk.Label(top, text="ENVIO AK2_GN PARA PROTHEUS", style="Title.TLabel").pack(anchor="w")

        file_fr = ttk.Frame(self)
        file_fr.pack(fill="x", padx=14, pady=(8, 4))
        ttk.Label(file_fr, text="Planilha:").pack(anchor="w")

        row1 = ttk.Frame(file_fr)
        row1.pack(fill="x", pady=(2, 0))
        self.entry_arquivo = ttk.Entry(row1, textvariable=self.arquivo_var)
        self.entry_arquivo.pack(side="left", fill="x", expand=True)
        ttk.Button(row1, text="PROCURAR", width=14, command=self._escolher_arquivo).pack(side="left", padx=(6, 0))
        self.btn_carregar = ttk.Button(row1, text="CARREGAR DADOS", width=18, command=self._carregar_planilha)
        self.btn_carregar.pack(side="left", padx=(6, 0))

        opt_fr = ttk.Frame(self)
        opt_fr.pack(fill="x", padx=14, pady=(6, 4))

        ttk.Label(opt_fr, text="AMBIENTE:").pack(side="left")
        ttk.Label(opt_fr, text="Produção (fixo)").pack(side="left", padx=(8, 18))

        ttk.Label(opt_fr, text="CENÁRIO:").pack(side="left")
        ttk.Radiobutton(
            opt_fr,
            text=MODO_BUDGET,
            value=MODO_BUDGET,
            variable=self.modo_var,
            command=self._on_modo_changed,
        ).pack(side="left", padx=(8, 0))
        ttk.Radiobutton(
            opt_fr,
            text=MODO_FORECAST,
            value=MODO_FORECAST,
            variable=self.modo_var,
            command=self._on_modo_changed,
        ).pack(side="left", padx=(8, 0))

        ttk.Button(
            opt_fr,
            text="▼ FILTRAR CENTROS DE CUSTO",
            width=28,
            command=self._abrir_seletor_centros,
        ).pack(side="left", padx=(20, 6))
        ttk.Button(
            opt_fr,
            text="LIMPAR FILTROS",
            width=16,
            command=self._limpar_selecao_centros,
        ).pack(side="left", padx=(0, 10))

        ttk.Button(
            opt_fr,
            text="LANÇAMENTO EXTRA ORÇAMENTO",
            width=28,
            command=self._abrir_extra_orcamento,
        ).pack(side="left")

        info_fr = ttk.Frame(self)
        info_fr.pack(fill="x", padx=14, pady=(2, 4))
        self.lbl_info_modo = ttk.Label(info_fr, text="Centros disponíveis: 0")
        self.lbl_info_modo.pack(side="left")
        self.lbl_info_sel = ttk.Label(info_fr, text="Selecionados: 0")
        self.lbl_info_sel.pack(side="left", padx=(20, 0))
        self.lbl_info_env = ttk.Label(info_fr, text="Ambiente atual: Produção")
        self.lbl_info_env.pack(side="left", padx=(20, 0))

        sel_fr = ttk.LabelFrame(self, text="Centro(s) de Custo selecionado(s) para envio")
        sel_fr.pack(fill="both", expand=False, padx=14, pady=(4, 8))
        self.lst_selecionados = Listbox(sel_fr, height=8)
        self.lst_selecionados.pack(fill="both", expand=True, padx=8, pady=8)

        status_fr = ttk.Frame(self)
        status_fr.pack(fill="x", padx=14, pady=(0, 4))
        ttk.Label(status_fr, textvariable=self.status_var).pack(anchor="w")

        btn_fr = ttk.Frame(self)
        btn_fr.pack(fill="x", padx=14, pady=(4, 8))
        self.btn_exec = ttk.Button(btn_fr, text="▶ ENVIAR DADOS", width=18, command=self._executar_envio)
        self.btn_exec.pack(side="left")
        ttk.Button(btn_fr, text="FECHAR", width=12, command=self.destroy).pack(side="right")

        log_fr = ttk.LabelFrame(self, text="Log")
        log_fr.pack(fill=BOTH, expand=True, padx=14, pady=(4, 12))
        self.logbox = ScrolledText(log_fr, height=18, font=("Consolas", 10))
        self.logbox.pack(fill=BOTH, expand=True, padx=6, pady=6)
        self.logbox.configure(state=DISABLED)

    def _set_busy(self, busy: bool):
        state = DISABLED if busy else NORMAL
        self.btn_exec.configure(state=state)
        self.btn_carregar.configure(state=state)

    def _append_log(self, msg: str):
        self.logbox.configure(state=NORMAL)
        self.logbox.insert(END, msg + ("\n" if not msg.endswith("\n") else ""))
        self.logbox.see(END)
        self.logbox.configure(state=DISABLED)

    def _poll_log(self):
        try:
            while True:
                msg = self.log_queue.get_nowait()
                self._append_log(msg)
        except Empty:
            pass
        self.after(120, self._poll_log)

    def _carregar_padrao_automaticamente(self):
        try:
            p = Path(self.arquivo_var.get().strip())
            if p.exists():
                self._carregar_planilha(auto=True)
            else:
                self.status_var.set("Planilha padrão não encontrada. Selecione outra planilha.")
        except Exception:
            self.status_var.set("Não foi possível carregar a planilha padrão automaticamente.")

    def _escolher_arquivo(self):
        inicial = str(self._arquivo_carregado.parent) if self._arquivo_carregado else str(DEFAULT_WORKBOOK.parent)
        arq = filedialog.askopenfilename(
            initialdir=inicial or None,
            title="Escolha a planilha",
            filetypes=[("Arquivos Excel", "*.xlsx")],
        )
        if arq:
            self.arquivo_var.set(arq)
            self._limpar_cache_arquivo()
            self._render_selecionados()
            self._atualizar_info_modo()

    def _limpar_cache_arquivo(self):
        self._arquivo_carregado = None
        self._cache = None
        self._centros_selecionados = []

    def _caminho_atual(self) -> Path:
        caminho = self.arquivo_var.get().strip()
        if not caminho:
            raise ValueError("Selecione a planilha que será lida.")
        p = Path(caminho)
        if not p.exists():
            raise FileNotFoundError(f"Arquivo não encontrado: {p}")
        if p.suffix.lower() != ".xlsx":
            raise ValueError("Selecione um arquivo .xlsx.")
        return p

    def _carregar_planilha(self, auto: bool = False):
        if self._worker and self._worker.is_alive():
            return

        try:
            workbook_path = self._caminho_atual()
        except Exception as e:
            if not auto:
                messagebox.showerror("Erro", str(e))
            else:
                self.status_var.set(str(e))
            return

        self.status_var.set("Validando rapidamente, criando backup e carregando dados...")
        self._set_busy(True)

        def run():
            try:
                logging.info("Iniciando validação rápida da planilha: %s", workbook_path)
                validar_estrutura_rapida(workbook_path)
                criar_backup(workbook_path)
                cache = carregar_dados_para_execucao(workbook_path)
                logging.info(
                    "Carga concluída | DADOS=%d | Centros Budget=%d | Centros Forecast=%d",
                    len(cache.dados),
                    len(cache.centros_por_modo[MODO_BUDGET]),
                    len(cache.centros_por_modo[MODO_FORECAST]),
                )

                def done():
                    self._arquivo_carregado = workbook_path
                    self._cache = cache
                    self._sincronizar_selecao_com_modo()
                    self._atualizar_info_modo()
                    self._render_selecionados()
                    self.status_var.set("Planilha carregada com sucesso.")
                    self._set_busy(False)
                    if not auto:
                        messagebox.showinfo("Planilha", "Planilha validada, backup criado e dados carregados.")

                self.after(0, done)
            except Exception as exc:
                logging.exception("Falha ao carregar planilha: %s", exc)

                def fail():
                    self.status_var.set("Falha ao carregar planilha.")
                    self._set_busy(False)
                    if not auto:
                        messagebox.showerror("Erro", str(exc))

                self.after(0, fail)

        self._worker = threading.Thread(target=run, daemon=True)
        self._worker.start()

    def _sincronizar_selecao_com_modo(self):
        if not self._cache:
            self._centros_selecionados = []
            return
        disponiveis = set(self._cache.centros_por_modo.get(self.modo_var.get(), []))
        self._centros_selecionados = [v for v in self._centros_selecionados if v in disponiveis]

    def _atualizar_info_modo(self):
        modo = self.modo_var.get()
        qtd = len(self._cache.centros_por_modo.get(modo, [])) if self._cache else 0
        self.lbl_info_modo.configure(text=f"Centros disponíveis no modo {modo}: {qtd}")
        self.lbl_info_sel.configure(text=f"Selecionados: {len(self._centros_selecionados)}")
        self.lbl_info_env.configure(text="Ambiente atual: Produção")

    def _render_selecionados(self):
        self.lst_selecionados.delete(0, END)
        if self._centros_selecionados:
            for item in self._centros_selecionados:
                self.lst_selecionados.insert(END, item)
        else:
            self.lst_selecionados.insert(END, "[Nenhum selecionado] -> serão enviados todos do modo atual.")
        self._atualizar_info_modo()

    def _on_modo_changed(self):
        self._sincronizar_selecao_com_modo()
        self._render_selecionados()

    def _abrir_seletor_centros(self):
        if not self._cache:
            messagebox.showwarning("Planilha", "Carregue a planilha antes de selecionar os centros de custo.")
            return

        modo = self.modo_var.get()
        disponiveis = self._cache.centros_por_modo.get(modo, [])
        empresas_map = self._cache.empresas_por_modo.get(modo, {})

        dlg = SelecionarCentrosDialog(
            self,
            valores=disponiveis,
            empresas_map=empresas_map,
            selecionados_iniciais=self._centros_selecionados,
        )
        self.wait_window(dlg)

        if dlg.resultado is not None:
            self._centros_selecionados = dlg.resultado
            self._render_selecionados()

    def _limpar_selecao_centros(self):
        self._centros_selecionados = []
        self._render_selecionados()

    def _abrir_extra_orcamento(self):
        ExtraOrcamentoDialog(self)

    def _executar_envio(self):
        if self._worker and self._worker.is_alive():
            return

        try:
            workbook_path = self._caminho_atual()
        except Exception as e:
            messagebox.showerror("Erro", str(e))
            return

        if not self._cache or self._arquivo_carregado != workbook_path:
            resposta = messagebox.askyesno(
                "Planilha não carregada",
                "A planilha ainda não foi carregada nesta sessão. Deseja carregar agora antes do envio?",
            )
            if not resposta:
                return
            self._carregar_planilha()
            return

        modo = self.modo_var.get()

        if not self._centros_selecionados:
            ok = messagebox.askyesno(
                "Enviar todos?",
                f"Nenhum Centro_Custo foi selecionado.\n\nDeseja enviar TODOS os registros do modo {modo} para Produção?"
            )
            if not ok:
                return

        self.status_var.set("Preparando envio para API de Produção...")
        self._set_busy(True)

        centros_selecionados = list(self._centros_selecionados)
        cache = self._cache

        def run():
            try:
                registros = preparar_registros_api(
                    dados=cache.dados,
                    mapa_emp_item=cache.mapa_emp_item,
                    mapa_item=cache.mapa_item,
                    modo=modo,
                    centros_selecionados=centros_selecionados,
                )

                if not registros:
                    raise ValueError("Nenhum registro elegível foi encontrado para envio.")

                qtd_lotes, qtd_registros = enviar_lotes_e_registrar(
                    workbook_path=workbook_path,
                    registros=registros,
                    modo=modo,
                )

                def done():
                    self.status_var.set("Envio concluído com sucesso.")
                    self._set_busy(False)
                    messagebox.showinfo(
                        "Envio concluído",
                        f"Ambiente: Produção\nLotes enviados: {qtd_lotes}\nRegistros enviados: {qtd_registros}\n\nLogs salvos em:\n{garantir_pasta_logs_api()}"
                    )

                self.after(0, done)
            except Exception as exc:
                logging.exception("Falha no envio: %s", exc)

                def fail():
                    self.status_var.set("Falha no envio.")
                    self._set_busy(False)
                    messagebox.showerror("Erro no envio", str(exc))

                self.after(0, fail)

        self._worker = threading.Thread(target=run, daemon=True)
        self._worker.start()


def main():
    app = App()
    app.mainloop()
    return 0


if __name__ == "__main__":
    sys.exit(main())